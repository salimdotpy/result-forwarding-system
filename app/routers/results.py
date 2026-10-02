import csv
import io
import uuid
from decimal import Decimal, InvalidOperation
from pathlib import Path
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import RedirectResponse
from openpyxl import load_workbook
from sqlmodel import select
from app.core.config import get_settings
from app.core.dependencies import SessionDep, current_user
from app.models import Course, Department, Faculty, ResultRecord, ResultSubmission, Student, User, WorkflowEvent
from app.models.enums import ResultStatus, Semester, UserRole, WorkflowAction
from app.services.audit import log_audit
from app.services.workflow import transition
from app.core.templates import templates, template_context

router = APIRouter(prefix="/results", tags=["results"])
settings = get_settings()


def visible_query(session, user: User):
    q = select(ResultSubmission)
    if user.role == UserRole.LECTURER:
        q = q.where(ResultSubmission.lecturer_id == user.id)
    elif user.role in {UserRole.EXAMINATION_OFFICER, UserRole.HOD}:
        q = q.where(ResultSubmission.department_id == user.department_id)
    elif user.role == UserRole.DEAN:
        q = q.where(ResultSubmission.faculty_id == user.faculty_id)
    return q


def can_view(submission: ResultSubmission, user: User) -> bool:
    if user.role == UserRole.ADMIN: return True
    if user.role == UserRole.LECTURER: return submission.lecturer_id == user.id
    if user.role in {UserRole.EXAMINATION_OFFICER, UserRole.HOD}: return submission.department_id == user.department_id
    if user.role == UserRole.DEAN: return submission.faculty_id == user.faculty_id
    return user.role == UserRole.ACADEMIC_PLANNING


def parse_rows(filename: str, data: bytes):
    suffix = Path(filename or "").suffix.lower()
    rows = []
    if suffix == ".csv":
        text = data.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            rows.append({k.strip().lower(): (v or "").strip() for k, v in row.items() if k})
    elif suffix == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values: return []
        headers = [str(v or "").strip().lower() for v in values[0]]
        for values_row in values[1:]:
            rows.append({headers[i]: str(values_row[i] or "").strip() for i in range(min(len(headers), len(values_row)))})
    else:
        raise HTTPException(422, "The uploaded file could not be processed. Please use CSV or XLSX format.")
    return rows


def normalize_row(row: dict, index: int):
    matric = row.get("student id") or row.get("matric number") or row.get("matric_number") or row.get("matric")
    name = row.get("student name") or row.get("student_name") or row.get("name")
    score = row.get("score") or row.get("result")
    grade = row.get("grade") or None
    remark = row.get("remark") or None
    if not matric or not name or score in (None, ""):
        raise ValueError(f"Row {index}: student ID, student name and score are required.")
    try:
        value = Decimal(str(score))
    except InvalidOperation as exc:
        raise ValueError(f"Row {index}: score must be numeric.") from exc
    if value < 0 or value > 100:
        raise ValueError(f"Row {index}: score must be between 0 and 100.")
    return str(matric), str(name), value, grade, remark

@router.get("")
async def list_results(request: Request, session: SessionDep, user: User = Depends(current_user), q: str = "", status_filter: str = ""):
    query = visible_query(session, user)
    if q:
        query = query.where(ResultSubmission.reference.contains(q) | ResultSubmission.academic_session.contains(q))
    if status_filter:
        query = query.where(ResultSubmission.status == status_filter)
    rows = session.exec(query.order_by(ResultSubmission.updated_at.desc())).all()
    return templates.TemplateResponse(request=request, name="results/list.html", context=template_context(request, user=user, rows=rows, q=q, status_filter=status_filter, statuses=[s.value for s in ResultStatus]))

@router.get("/new")
async def new_result(request: Request, session: SessionDep, user: User = Depends(current_user)):
    if user.role != UserRole.LECTURER:
        raise HTTPException(403, "You do not have permission to perform this action.")
    faculties = session.exec(select(Faculty)).all()
    departments = session.exec(select(Department)).all()
    courses = session.exec(select(Course)).all()
    return templates.TemplateResponse(request=request, name="results/new.html", context=template_context(request, user=user, faculties=faculties, departments=departments, courses=courses, semesters=list(Semester)))

@router.post("/new")
async def create_result(
    request: Request, session: SessionDep, user: User = Depends(current_user),
    academic_session: str = Form(...), semester: Semester = Form(...), faculty_id: int = Form(...), department_id: int = Form(...), course_id: int = Form(...), file: UploadFile = File(...),
):
    if user.role != UserRole.LECTURER:
        raise HTTPException(403, "You do not have permission to perform this action.")
    data = await file.read()
    if len(data) > settings.max_upload_size:
        return templates.TemplateResponse("results/new.html", template_context(request, user=user, error=f"The uploaded file is too large. Maximum size is {settings.max_upload_size_mb} MB."), status_code=422)
    try:
        parsed = parse_rows(file.filename or "", data)
        if not parsed:
            raise ValueError("The uploaded file contains no result records.")
        normalized = [normalize_row(row, i) for i, row in enumerate(parsed, 2)]
    except (UnicodeDecodeError, ValueError, HTTPException) as exc:
        message = exc.detail if isinstance(exc, HTTPException) else str(exc)
        faculties = session.exec(select(Faculty)).all(); departments = session.exec(select(Department)).all(); courses = session.exec(select(Course)).all()
        return templates.TemplateResponse(request=request, name="results/new.html", context=template_context(request, user=user, faculties=faculties, departments=departments, courses=courses, semesters=list(Semester), error=message), status_code=422)

    department = session.get(Department, department_id)
    course = session.get(Course, course_id)
    faculty = session.get(Faculty, faculty_id)
    if not department or not course or not faculty or department.faculty_id != faculty_id or course.department_id != department_id:
        raise HTTPException(422, "Please select a valid faculty, department and course combination.")

    submission = ResultSubmission(reference=f"RES-{uuid.uuid4().hex[:8].upper()}", academic_session=academic_session.strip(), semester=semester, faculty_id=faculty_id, department_id=department_id, course_id=course_id, lecturer_id=user.id, status=ResultStatus.DRAFT)
    session.add(submission); session.flush()
    for matric, name, score, grade, remark in normalized:
        student = session.exec(select(Student).where(Student.matric_number == matric)).first()
        session.add(ResultRecord(submission_id=submission.id, student_id=student.id if student else None, matric_number=matric, student_name=name, score=score, grade=grade, remark=remark))
    session.commit(); session.refresh(submission)
    log_audit(session, user=user, action=WorkflowAction.UPLOADED.value, submission_id=submission.id, previous_status=ResultStatus.DRAFT.value, new_status=ResultStatus.DRAFT.value)
    transition(session, submission, user, ResultStatus.SUBMITTED, WorkflowAction.SUBMITTED)
    session.commit()
    return RedirectResponse(f"/results/{submission.id}", status_code=303)

@router.get("/{submission_id}")
async def detail(request: Request, submission_id: int, session: SessionDep, user: User = Depends(current_user)):
    submission = session.get(ResultSubmission, submission_id)
    if not submission or not can_view(submission, user):
        raise HTTPException(404, "The requested result could not be found.")
    records = session.exec(select(ResultRecord).where(ResultRecord.submission_id == submission.id).order_by(ResultRecord.matric_number)).all()
    events = session.exec(select(WorkflowEvent).where(WorkflowEvent.submission_id == submission.id).order_by(WorkflowEvent.created_at)).all()
    actor_names = {e.actor_id: session.get(User, e.actor_id).full_name for e in events if session.get(User, e.actor_id)}
    course = session.get(Course, submission.course_id); department = session.get(Department, submission.department_id); faculty = session.get(Faculty, submission.faculty_id); lecturer = session.get(User, submission.lecturer_id)
    return templates.TemplateResponse(request=request, name="results/detail.html", context=template_context(request, user=user, submission=submission, records=records, events=events, actor_names=actor_names, course=course, department=department, faculty=faculty, lecturer=lecturer))

@router.post("/{submission_id}/action")
async def action(request: Request, submission_id: int, session: SessionDep, user: User = Depends(current_user), action_name: str = Form(...), comment: str = Form("")):
    submission = session.get(ResultSubmission, submission_id)
    if not submission or not can_view(submission, user):
        raise HTTPException(404, "The requested result could not be found.")
    mapping = {
        "start_verification": (ResultStatus.UNDER_VERIFICATION, WorkflowAction.START_VERIFICATION),
        "verify": (ResultStatus.VERIFIED, WorkflowAction.VERIFIED),
        "return": (ResultStatus.RETURNED, WorkflowAction.RETURNED),
        "reject": (ResultStatus.REJECTED, WorkflowAction.REJECTED),
        "resubmit": (ResultStatus.SUBMITTED, WorkflowAction.SUBMITTED),
        "hod_approve": (ResultStatus.HOD_APPROVED, WorkflowAction.HOD_APPROVED),
        "dean_approve": (ResultStatus.DEAN_APPROVED, WorkflowAction.DEAN_APPROVED),
        "forward": (ResultStatus.FORWARDED, WorkflowAction.FORWARDED),
        "receive": (ResultStatus.RECEIVED, WorkflowAction.RECEIVED),
    }
    if action_name not in mapping:
        raise HTTPException(422, "This action is not available.")
    new_status, workflow_action = mapping[action_name]
    transition(session, submission, user, new_status, workflow_action, comment.strip() or None)
    session.commit()
    return RedirectResponse(f"/results/{submission.id}", status_code=303)
